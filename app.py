"""SettlementCheck

Potrebno: streamlit>=1.40, pandas, requests, rapidfuzz
Secrets (.streamlit/secrets.toml):
    APP_PASSWORD = "nova-jaka-lozinka"
    APISPORTS_KEY = "tvoj-api-kljuc"        # opciono
    FOOTBALLDATA_KEY = "tvoj-fd-kljuc"     # opciono
    RAPIDAPI_KEY = "tvoj-rapidapi-kljuc"    # opciono (SportAPI)
    DEFAULT_PROVIDER = "SportAPI"           # opciono
"""
import hmac
import io
import json
import math
import os
import re
import struct
import time
import unicodedata
import wave
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import pandas as pd
import requests
import streamlit as st
from rapidfuzz import fuzz

st.set_page_config(page_title="SettlementCheck", page_icon=":material/fact_check:", layout="wide")

# ==========================================
# 0. KONSTANTE
# ==========================================
STATE_FILE = os.environ.get("SETTLEMENT_STATE_FILE", "settlement_state.json")
API_URL = "https://v3.football.api-sports.io/fixtures"
FD_URL = "https://api.football-data.org/v4/matches"

# SportAPI konfiguracija (RapidAPI - sportapi7.p.rapidapi.com)
SPORTAPI_HOST = "sportapi7.p.rapidapi.com"
SPORTAPI_BASE = "https://sportapi7.p.rapidapi.com/api/v1"

TZ = ZoneInfo("Europe/Belgrade")
MATCH_THRESHOLD = 80  # minimalna sličnost imena timova (0-100)
PROVIDERS = ["API-Sports", "football-data.org", "SportAPI"]

FD_STATUS = {
    "IN_PLAY": "LIVE", "LIVE": "LIVE", "PAUSED": "HT", "FINISHED": "FT", "AWARDED": "FT",
    "SUSPENDED": "SUSP", "INTERRUPTED": "SUSP", "POSTPONED": "PST", "CANCELLED": "CANC",
    "CANCELED": "CANC"
}

SPORTAPI_STATUS = {
    "notstarted": "NS",
    "in-progress": "LIVE",
    "finished": "FT",
    "postponed": "PST",
    "canceled": "CANC",
    "delayed": "PST",
    "interrupted": "SUSP",
    "suspended": "SUSP"
}


def ov_key(provider, mid):
    """Ručna spajanja su vezana za provajdera."""
    return mid if provider == "API-Sports" else f"{provider}|{mid}"

FINISHED = {"FT", "AET", "PEN"}
SKIP_STATUSES = {"NS", "TBD", "PST", "CANC"}
QUALIFIERS = {"u17", "u18", "u19", "u20", "u21", "u23", "w", "women", "ii", "b", "reserves"}
IGNORE_TOKENS = {"A", "M"}
BASE_ALIASES = {
    "red star": "crvena zvezda",
    "man utd": "manchester united",
    "man city": "manchester city",
    "partizan": "partizan beograd",
    "inter": "inter milan",
    "ac milan": "milan",
}
PING = [(1320, 0.18)]
STATUS_RANK = {"Neslaganje": 0, "Čeka potvrdu": 1, "Nema podataka": 2,
               "Sistem bez rezultata": 3, "OK": 4}


def get_secret(name, default=""):
    try:
        return st.secrets[name]
    except Exception:
        return default


def g(d, *path, default=None):
    """Bezbedno čitanje ugnježdenih ključeva."""
    for k in path:
        if not isinstance(d, dict):
            return default
        d = d.get(k)
        if d is None:
            return default
    return d


# ==========================================
# DIZAJN (crna + zelena)
# ==========================================
CSS = """
@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;500;600&display=swap');
:root{--g:#00A651;--gd:#007A3D;--line:#D9E0DC;--ink:#0D0F0E;}
p,label,input,textarea,button,h1,h2,h3,h4,li,td,th,
[data-testid="stMetricValue"],[data-testid="stMetricLabel"],[data-testid="stCaptionContainer"]
{font-family:'IBM Plex Sans',system-ui,-apple-system,'Segoe UI',sans-serif;}
.block-container{padding-top:1.4rem;max-width:1400px;}
footer{visibility:hidden;}
[data-testid="stHeader"]{background:transparent;}
h1,h2,h3{font-weight:600;letter-spacing:-.01em;}

.sc-header{display:flex;align-items:center;gap:14px;background:#000;border-radius:8px;padding:14px 18px;margin-bottom:16px;}
.sc-mark{width:34px;height:34px;border-radius:7px;background:var(--g);display:grid;place-items:center;flex:none;}
.sc-name{color:#fff;font-size:1.25rem;font-weight:600;letter-spacing:-.01em;line-height:1.2;}
.sc-sub{color:#8FA39A;font-size:.82rem;}

.sc-status{background:#000;border-left:8px solid var(--c);border-radius:6px;padding:14px 18px;margin:0 0 14px;}
.sc-st-title{color:#fff;font-size:1.5rem;font-weight:600;line-height:1.25;font-variant-numeric:tabular-nums;}
.sc-st-sub{color:#9AA8A1;font-size:.88rem;margin-top:2px;}

[data-testid="stMetric"]{background:#fff;border:1px solid var(--line);border-radius:6px;padding:12px 14px;}
[data-testid="stMetricLabel"] p{font-size:.78rem;color:#5B6761;}
[data-testid="stMetricValue"]{font-size:1.7rem;font-weight:600;font-variant-numeric:tabular-nums;}

.stTabs [data-baseweb="tab-list"]{gap:4px;border-bottom:1px solid var(--line);}
.stTabs [data-baseweb="tab"]{padding:8px 14px;font-weight:500;}
.stTabs [aria-selected="true"]{color:#000;font-weight:600;}
.stTabs [data-baseweb="tab-highlight"]{background:var(--g);height:3px;}

.stButton>button,.stDownloadButton>button{border-radius:6px;font-weight:500;}
button[kind="primary"],[data-testid="stBaseButton-primary"]{background:var(--g);border:0;color:#000;}
button[kind="primary"]:hover,[data-testid="stBaseButton-primary"]:hover{background:#00BD5B;color:#000;}
button[kind="primary"] p,[data-testid="stBaseButton-primary"] p{color:#000;font-weight:600;}
[data-testid="stExpander"]{background:#fff;border:1px solid var(--line);border-radius:6px;}
[data-testid="stAlert"]{border-radius:6px;}

[data-testid="stSidebar"]{background:#000;border-right:1px solid #1c1c1c;}
[data-testid="stSidebar"] p,[data-testid="stSidebar"] label,[data-testid="stSidebar"] span,
[data-testid="stSidebar"] li{color:#E3E9E6;}
[data-testid="stSidebar"] h2{color:#fff;font-size:1.05rem;}
[data-testid="stSidebar"] h3{color:#7FE0A9;font-size:.85rem;font-weight:600;margin-top:.6rem;}
[data-testid="stSidebar"] input,[data-testid="stSidebar"] textarea,
[data-testid="stSidebar"] [data-baseweb="select"]>div,[data-testid="stSidebar"] [data-baseweb="input"],
[data-testid="stSidebar"] [data-baseweb="base-input"]{background:#141414 !important;color:#fff !important;border-color:#2b2b2b !important;}
[data-testid="stSidebar"] .stButton>button{background:#141414;color:#fff;border:1px solid #333;}
[data-testid="stSidebar"] .stButton>button:hover{border-color:var(--g);color:var(--g);}
[data-testid="stSidebar"] [data-testid="stAlert"] p,[data-testid="stSidebar"] [data-testid="stAlert"] span{color:#1a1a1a;}
.sc-req{background:#2a0b0b;border-left:4px solid #E5393A;color:#fff;padding:8px 10px;border-radius:4px;font-size:.85rem;margin:-6px 0 10px;}
"""

LOGO_SVG = ('<svg width="20" height="20" viewBox="0 0 24 24" fill="none" stroke="#000" stroke-width="2.8" '
            'stroke-linecap="round" stroke-linejoin="round"><path d="M5 12.5l4.5 4.5L19 7.5"/></svg>')


def inject_css():
    st.markdown(f"<style>{CSS}</style>", unsafe_allow_html=True)


def app_header():
    st.markdown(
        f'<div class="sc-header"><div class="sc-mark">{LOGO_SVG}</div><div>'
        '<div class="sc-name">SettlementCheck</div><div class="sc-sub">Provera rezultata uživo</div>'
        '</div></div>', unsafe_allow_html=True)


inject_css()


# ==========================================
# 1. AUTORIZACIJA
# ==========================================
def check_password():
    if st.session_state.get("authenticated"):
        return True
    expected = str(get_secret("APP_PASSWORD"))
    if not expected:
        st.error("APP_PASSWORD nije podešen u .streamlit/secrets.toml")
        return False
    app_header()
    with st.form("login"):
        pwd = st.text_input("Lozinka:", type="password")
        ok = st.form_submit_button("Prijavi se")
    if ok:
        if hmac.compare_digest(pwd.encode(), expected.encode()):
            st.session_state["authenticated"] = True
            st.rerun()
        time.sleep(1.5)
        st.error("Netačna lozinka!")
    return False


if not check_password():
    st.stop()


# ==========================================
# 2. TRAJNO STANJE
# ==========================================
def load_store():
    try:
        with open(STATE_FILE, encoding="utf-8") as f:
            d = json.load(f)
    except Exception:
        d = {}
    return {k: d.get(k, {}) for k in ("overrides", "aliases", "incidents")}


def save_store():
    try:
        tmp = STATE_FILE + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(st.session_state["store"], f, ensure_ascii=False)
        os.replace(tmp, STATE_FILE)
    except Exception as e:
        st.toast(f"Čuvanje stanja nije uspelo: {e}")


if "store" not in st.session_state:
    st.session_state["store"] = load_store()
for _k, _v in {"process_triggered": False, "refresh_nonce": 0, "pending": {}}.items():
    st.session_state.setdefault(_k, _v)


# ==========================================
# 3. ZVUK
# ==========================================
@st.cache_resource
def make_wav():
    rate, frames = 22050, bytearray()
    for freq, dur in PING:
        for i in range(int(rate * dur)):
            v = int(12000 * math.sin(2 * math.pi * freq * i / rate)) if freq else 0
            frames += struct.pack("<h", v)
    buf = io.BytesIO()
    with wave.open(buf, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(rate)
        w.writeframes(bytes(frames))
    return buf.getvalue()


# ==========================================
# 4. NORMALIZACIJA I POREĐENJE TIMOVA
# ==========================================
def strip_accents(s):
    s = s.replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def clean_name(name):
    n = strip_accents(name).lower()
    n = re.sub(r"[^\w\s]", " ", n)
    n = re.sub(r"\b(fc|fk|cd|sc|club)\b", " ", n)
    return re.sub(r"\s+", " ", n).strip()


def normalize_team_name(name):
    c = clean_name(name)
    return st.session_state["store"]["aliases"].get(c) or BASE_ALIASES.get(c, c)


def same_qualifiers(a, b):
    return (set(a.split()) & QUALIFIERS) == (set(b.split()) & QUALIFIERS)


def find_match(nh, na, index, threshold):
    scored = []
    for fx in index:
        if not (same_qualifiers(nh, fx["nh"]) and same_qualifiers(na, fx["na"])):
            continue
        sh = fuzz.token_sort_ratio(nh, fx["nh"])
        sa = fuzz.token_sort_ratio(na, fx["na"])
        if sh >= threshold and sa >= threshold:
            scored.append(((sh + sa) / 2, fx))
    if not scored:
        return None, 0, False
    scored.sort(key=lambda t: t[0], reverse=True)
    ambiguous = len(scored) > 1 and scored[0][0] - scored[1][0] < 3
    return scored[0][1], scored[0][0], ambiguous


def find_candidate(m, index):
    nh, na = normalize_team_name(m["Home"]), normalize_team_name(m["Away"])
    best, best_s = None, 60
    for fx in index:
        s = (fuzz.token_sort_ratio(nh, fx["nh"]) + fuzz.token_sort_ratio(na, fx["na"])) / 2
        if s > best_s:
            best, best_s = fx, s
    return best


def apply_override(m, fx):
    store = st.session_state["store"]
    store["overrides"][ov_key(st.session_state.get("provider", "API-Sports"), m["ID"])] = {
        "fixture_id": fx["id"], "label": f"{fx['home']} - {fx['away']}"}
    for mine, theirs in ((m["Home"], fx["home"]), (m["Away"], fx["away"])):
        a, b = clean_name(mine), clean_name(theirs)
        if a and b and a != b and mine not in ("Domaćin", "Gost"):
            store["aliases"][a] = b
    save_store()


# ==========================================
# 5. PARSER KOPIRANOG TEKSTA
# ==========================================
def parse_copied_data(text):
    text = text.replace("\t", "\n").replace("\r", "")
    lines = [re.sub(r"\s+", " ", ln).strip() for ln in text.split("\n") if ln.strip()]
    id_pos = [i for i, ln in enumerate(lines) if re.match(r"^\d{4,6}$", ln)]
    matches = []

    for n, idx in enumerate(id_pos):
        end = min(idx + 18, len(lines))
        if n + 1 < len(id_pos):
            end = min(end, max(id_pos[n + 1] - 1, idx + 1))
        block = lines[idx + 1:end]
        league = lines[idx - 1] if idx > 0 else "Nepoznata liga"

        scores, ht, teams = [], "N/A", []
        for item in block:
            if re.match(r"^\d+:\d+$", item):
                ht = item
            elif item.isdigit() and len(item) <= 2:
                scores.append(item)
            elif (not item.isdigit() and item not in IGNORE_TOKENS and ":" not in item
                  and len(item) > 1 and item != league):
                teams.append(item)

        home = teams[0] if len(teams) > 0 else "Domaćin"
        away = teams[1] if len(teams) > 1 else "Gost"
        matches.append({
            "ID": lines[idx],
            "Liga": league,
            "Home": home,
            "Away": away,
            "Meč": f"{home} - {away}",
            "Moj Sistem Rezultat": f"{scores[0]}:{scores[1]}" if len(scores) >= 2 else "N/A",
            "Score2 (HT)": ht,
            "Oznaka": "A" if "A" in block else "M",
        })
    return matches


# ==========================================
# 6. API POZIVI
# ==========================================
def _request(api_key, params):
    r = requests.get(API_URL, headers={"x-apisports-key": api_key}, params=params, timeout=8)
    r.raise_for_status()
    body = r.json()
    if body.get("errors"):
        raise RuntimeError(str(body["errors"]))
    return (body.get("response", []), r.headers.get("x-ratelimit-requests-remaining", "N/A"),
            r.headers.get("x-ratelimit-requests-limit", "N/A"), time.time())


def fd_to_fixture(m):
    sc = m.get("score") or {}
    ft, reg, ex = sc.get("fullTime") or {}, sc.get("regularTime") or {}, sc.get("extraTime") or {}
    duration = sc.get("duration")
    short = FD_STATUS.get(m.get("status"), "NS")
    if short == "FT" and duration == "EXTRA_TIME":
        short = "AET"
    elif short == "FT" and duration == "PENALTY_SHOOTOUT":
        short = "PEN"
    gh, ga = ft.get("home"), ft.get("away")
    if duration in ("EXTRA_TIME", "PENALTY_SHOOTOUT") and reg.get("home") is not None:
        gh, ga = reg["home"] + (ex.get("home") or 0), reg["away"] + (ex.get("away") or 0)
    return {
        "fixture": {"id": m.get("id"), "status": {"short": short, "elapsed": m.get("minute")}},
        "league": {"name": g(m, "competition", "name", default="")},
        "teams": {"home": {"name": g(m, "homeTeam", "name", default="")},
                  "away": {"name": g(m, "awayTeam", "name", default="")}},
        "goals": {"home": gh, "away": ga},
        "score": {"fulltime": {"home": reg.get("home", gh), "away": reg.get("away", ga)}},
    }


def _request_fd(api_key, params):
    r = requests.get(FD_URL, headers={"X-Auth-Token": api_key}, params=params, timeout=8)
    r.raise_for_status()
    matches = [fd_to_fixture(m) for m in r.json().get("matches", [])]
    return matches, r.headers.get("X-Requests-Available-Minute", "N/A"), "N/A", time.time()


def sportapi_to_fixture(event):
    """Prevodi SportAPI strukturu meča u standardizovan format."""
    status_type = g(event, "status", "type", default="notstarted")
    short = SPORTAPI_STATUS.get(status_type, "NS")
    
    home_score = g(event, "homeScore", "current", default=0)
    away_score = g(event, "awayScore", "current", default=0)
    
    ht_home = g(event, "homeScore", "period1")
    ht_away = g(event, "awayScore", "period1")
    
    return {
        "fixture": {"id": event.get("id"), "status": {"short": short, "elapsed": None}},
        "league": {"name": g(event, "tournament", "name", default="")},
        "teams": {"home": {"name": g(event, "homeTeam", "name", default="")},
                  "away": {"name": g(event, "awayTeam", "name", default="")}},
        "goals": {"home": home_score, "away": away_score},
        "score": {"fulltime": {"home": home_score, "away": away_score},
                  "halftime": {"home": ht_home, "away": ht_away}},
    }


def _request_sportapi(api_key, is_live=True, day=None):
    headers = {
        "x-rapidapi-key": api_key,
        "x-rapidapi-host": SPORTAPI_HOST
    }
    
    if is_live:
        url = f"{SPORTAPI_BASE}/sport/football/events/live"
    else:
        req_day = day or date.today().isoformat()
        url = f"{SPORTAPI_BASE}/sport/football/scheduled-events/{req_day}"
        
    r = requests.get(url, headers=headers, timeout=8)
    r.raise_for_status()
    data = r.json()
    events = data.get("events", [])
    matches = [sportapi_to_fixture(ev) for ev in events]
    
    rem = r.headers.get("x-ratelimit-requests-remaining", "N/A")
    lim = r.headers.get("x-ratelimit-requests-limit", "N/A")
    return matches, rem, lim, time.time()


@st.cache_data(ttl=15, show_spinner=False)
def _fetch_live(provider, api_key, nonce):
    if provider == "football-data.org":
        return _request_fd(api_key, {"status": "LIVE"})
    elif provider == "SportAPI":
        return _request_sportapi(api_key, is_live=True)
    return _request(api_key, {"live": "all"})


@st.cache_data(ttl=120, show_spinner=False)
def _fetch_date(provider, api_key, day, nonce):
    if provider == "football-data.org":
        return _request_fd(api_key, {"dateFrom": day, "dateTo": (date.fromisoformat(day) + timedelta(days=1)).isoformat()})
    elif provider == "SportAPI":
        return _request_sportapi(api_key, is_live=False, day=day)
    return _request(api_key, {"date": day, "timezone": "Europe/Belgrade"})


def load_fixtures(provider, api_key, include_finished, day, nonce, reuse=False):
    key = (provider, hash(api_key), include_finished, day, nonce)
    saved = st.session_state.get("feed_cache")
    if reuse and saved and saved["key"] == key:
        return saved["value"]
    merged, remaining, limit, errors, stamps, live_ts = {}, "N/A", "N/A", [], [], None
    calls = [("live", lambda: _fetch_live(provider, api_key, nonce))]
    if include_finished:
        calls.insert(0, ("datum", lambda: _fetch_date(provider, api_key, day, nonce)))
    for name, fn in calls:
        try:
            data, remaining, limit, ts = fn()
            stamps.append(ts)
            if name == "live":
                live_ts = ts
            for fx in data:
                merged[g(fx, "fixture", "id")] = fx
        except requests.HTTPError as e:
            code = e.response.status_code if e.response is not None else "?"
            errors.append("Prekoračen limit API zahteva (sačekajte ili proverite plan)." if code == 429 else
                          "Ključ nije prihvaćen ili plan ne pokriva zahtev (HTTP %s)." % code if code in (401, 403)
                          else f"{name}: HTTP {code}")
        except Exception as e:
            errors.append(f"{name}: {e}")
    result = (list(merged.values()), remaining, limit, live_ts or (min(stamps) if stamps else None), errors)
    if not errors:
        st.session_state["feed_cache"] = {"key": key, "value": result}
    return result


def build_index(fixtures, use_90):
    idx = []
    for fx in fixtures:
        short = g(fx, "fixture", "status", "short", default="")
        if short in SKIP_STATUSES:
            continue
        if use_90 and short in {"AET", "PEN"}:
            gh, ga = g(fx, "score", "fulltime", "home", default=0), g(fx, "score", "fulltime", "away", default=0)
        else:
            gh, ga = g(fx, "goals", "home", default=0), g(fx, "goals", "away", default=0)
        elapsed = g(fx, "fixture", "status", "elapsed")
        home, away = g(fx, "teams", "home", "name", default=""), g(fx, "teams", "away", "name", default="")
        idx.append({
            "id": g(fx, "fixture", "id"),
            "liga": g(fx, "league", "name", default=""),
            "home": home, "away": away,
            "nh": normalize_team_name(home), "na": normalize_team_name(away),
            "score": f"{gh}:{ga}",
            "short": short,
            "status": f"{short} {elapsed}'" if elapsed and short not in FINISHED else short,
        })
    return idx


# ==========================================
# 7. INCIDENTI
# ==========================================
def clear_resolved():
    store = st.session_state["store"]
    store["incidents"] = {k: v for k, v in store["incidents"].items() if v["stanje"] != "rešeno"}
    save_store()


def clear_overrides():
    st.session_state["store"]["overrides"] = {}
    save_store()


def highlight_status(val):
    return {
        "Neslaganje": "background-color:#FDE3E3;color:#9B1C1C;font-weight:600;",
        "Čeka potvrdu": "background-color:#FFF1CC;color:#7A5200;font-weight:600;",
        "Nema podataka": "background-color:#E9EDEB;color:#4A5550;",
        "Sistem bez rezultata": "background-color:#ECE4F6;color:#53257F;",
        "OK": "background-color:#DDF3E6;color:#0A5A2C;font-weight:600;",
    }.get(val, "")


RESULT_STYLE = {
    "Neslaganje": "color:#C62828;font-weight:700;",
    "Čeka potvrdu": "color:#B26A00;font-weight:700;",
}
DISPLAY_COLS = ["Status", "Incident", "ID", "Oznaka", "Liga", "Meč", "Rezultat",
                "Score2 (HT)", "Status Meča", "Pouzdanost"]


def status_banner(n_active, n_wait, n_na, n_ok, feed_problem=""):
    if n_active:
        color, title = "#E5393A", f"{n_active} {'neslaganje' if n_active == 1 else 'neslaganja'}"
        sub = f"Čeka potvrdu: {n_wait}. Bez podataka: {n_na}."
        if feed_problem:
            sub += f" Upozorenje feeda: {feed_problem}."
    elif feed_problem:
        color, title = "#8A8F8C", "Provera nije pouzdana"
        sub = f"{feed_problem[0].upper() + feed_problem[1:]}. Rezultati ispod mogu biti netačni."
    elif n_wait:
        color, title = "#F2A900", f"Čeka potvrdu: {n_wait}"
        sub = "Moguće kašnjenje feeda, proverava se ponovo."
    else:
        color, title = "#00A651", "Nema neslaganja"
        sub = f"Usklađeno: {n_ok}." + (f" Bez podataka: {n_na}." if n_na else "")
    st.markdown(f'<div class="sc-status" style="--c:{color}"><div class="sc-st-title">{title}</div>'
                f'<div class="sc-st-sub">{sub}</div></div>', unsafe_allow_html=True)


def style_conf(v):
    if v >= 95:
        return "background-color:#DDF3E6;color:#0A5A2C;"
    if v >= 80:
        return "background-color:#FFF1CC;color:#7A5200;"
    return "background-color:#FFE1CC;color:#8A3B00;" if v > 0 else ""


def style_row(row):
    css = RESULT_STYLE.get(row["Status"], "")
    return [css if c == "Rezultat" else "" for c in row.index]


def render_table(d):
    if d.empty:
        st.caption("Nema mečeva u ovoj kategoriji.")
        return
    stl = d[DISPLAY_COLS].style
    mp = stl.map if hasattr(stl, "map") else stl.applymap
    stl = mp(highlight_status, subset=["Status"])
    stl = mp(style_conf, subset=["Pouzdanost"])
    stl = stl.apply(style_row, axis=1)
    st.dataframe(stl, use_container_width=True, hide_index=True)


def stamp():
    return datetime.now(TZ).strftime("%d.%m %H:%M:%S")


def fmt_time(ts):
    return datetime.fromtimestamp(ts, TZ).strftime("%H:%M:%S")


def fmt_dur(sec):
    sec = max(int(sec), 0)
    return f"{sec // 3600} h {sec % 3600 // 60} min" if sec >= 3600 else f"{sec // 60} min {sec % 60} s"


INC_COLS = ["Otkriveno", "ID", "Liga", "Meč", "Tvoj Sistem", "Live", "Stanje", "Operater", "Rešeno", "Trajanje"]
EXCEL_COLS = ["Status", "ID", "Oznaka", "Liga", "Meč", "Tvoj Sistem", "Teren / Live Feed",
              "Score2 (HT)", "Status Meča", "Pouzdanost"]
XL_FILL = {"Neslaganje": "FDE3E3", "Čeka potvrdu": "FFF1CC", "Nema podataka": "E9EDEB",
           "Sistem bez rezultata": "ECE4F6", "OK": "DDF3E6"}


def incidents_frame(inc_store, now):
    rows = []
    for v in inc_store.values():
        end = v.get("razreseno_ts") or (None if v["stanje"] == "rešeno" else now)
        rows.append({
            "Otkriveno": v["otkriveno"], "ID": v["id"], "Liga": v["liga"], "Meč": v["meč"],
            "Tvoj Sistem": v["sistem"], "Live": v["live"], "Stanje": v["stanje"],
            "Operater": v.get("operater", ""), "Rešeno": v.get("razreseno", ""),
            "Trajanje": fmt_dur(end - v["ts"]) if v.get("ts") and end else ""})
    return pd.DataFrame(rows, columns=INC_COLS)


def build_excel(df_full, inc_store, meta):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    lg = df_full.groupby("Liga")["Status"].value_counts().unstack(fill_value=0)
    lg = lg.reindex(columns=list(STATUS_RANK), fill_value=0)
    lg["Ukupno"] = lg.sum(axis=1)
    lg = lg.sort_values(["Neslaganje", "Nema podataka"], ascending=False).reset_index()
    sheets = {
        "Pregled": df_full[EXCEL_COLS],
        "Neslaganja": incidents_frame(inc_store, time.time()),
        "Po ligama": lg,
        "Info": pd.DataFrame(meta, columns=["Stavka", "Vrednost"]),
    }
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, d in sheets.items():
            d.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="000000")
                c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            if name != "Info":
                ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(d.columns, 1):
                width = max([len(str(col))] + [len(str(x)) for x in d[col].head(500)]) + 3
                ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
            if "Status" in d.columns:
                ci = list(d.columns).index("Status") + 1
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row, column=ci)
                    if cell.value in XL_FILL:
                        cell.fill = PatternFill("solid", fgColor=XL_FILL[cell.value])
    return buf.getvalue()


# ==========================================
# 8. SIDEBAR
# ==========================================
app_header()
st.sidebar.text_input("Ime operatera (obavezno)", key="operator", placeholder="Ime i prezime")
operator_ok = bool(st.session_state.get("operator", "").strip())
if not operator_ok:
    st.sidebar.markdown('<div class="sc-req">Obavezno: unesite ime operatera. Bez njega se provera ne pokreće.</div>',
                        unsafe_allow_html=True)
st.sidebar.header("Podešavanja")

default_provider = get_secret("DEFAULT_PROVIDER", "API-Sports")
def_idx = 0
if default_provider == "football-data.org":
    def_idx = 1
elif default_provider in ("SportAPI", "RapidAPI"):
    def_idx = 2

provider = st.sidebar.selectbox("Provajder podataka:", PROVIDERS, key="provider", index=def_idx)

if st.session_state.get("last_provider") != provider:
    for _k in ("api_remaining", "api_limit", "feed_cache"):
        st.session_state.pop(_k, None)
    st.session_state["last_provider"] = provider

if provider == "football-data.org":
    secret_name = "FOOTBALLDATA_KEY"
elif provider == "SportAPI":
    secret_name = "RAPIDAPI_KEY"
else:
    secret_name = "APISPORTS_KEY"

api_key_input = st.sidebar.text_input(f"{provider} ključ (opciono, zamenjuje secrets):", type="password",
                                      key=f"key_{provider}")
api_key = (api_key_input or str(get_secret(secret_name))).strip()

if provider == "football-data.org":
    st.sidebar.caption("Besplatan plan: 10 zahteva u minuti i ograničen broj takmičenja.")
elif provider == "SportAPI":
    st.sidebar.caption("SportAPI preko RapidAPI platforme (sportapi7.p.rapidapi.com).")

st.sidebar.subheader("Izvor podataka")
include_finished = st.sidebar.checkbox("Uključi i završene mečeve (po datumu)", value=True)
day = st.sidebar.date_input("Datum:", value=date.today()).isoformat() if include_finished else ""
use_90 = st.sidebar.checkbox("Za AET/PEN koristi rezultat posle 90 min", value=False)

st.sidebar.subheader("Tolerancija kašnjenja")
grace = st.sidebar.slider("Neslaganje je alarm tek posle (s):", 0, 300, 60, 15,
                          help="Završeni mečevi (FT/AET/PEN) se ne odlažu.")

st.sidebar.subheader("Osvežavanje")
refresh_mode = st.sidebar.radio("Režim:", ["Manuelno (Ručno)", "Automatsko"], index=0)
auto_interval = 60
if refresh_mode == "Automatsko":
    auto_interval = st.sidebar.slider("Interval (sekunde):", 15, 300, 60, 15)
    st.sidebar.warning("Troši do 2 API zahteva po ciklusu.")

st.sidebar.subheader("Zvučna upozorenja")
sound_alert = st.sidebar.checkbox("Zvučni ping za nova neslaganja", value=True)

if st.sidebar.button("Odjavi se", icon=":material/logout:"):
    st.session_state["authenticated"] = False
    st.rerun()


# ==========================================
# 9. UNOS TEKSTA
# ==========================================
def clear_text_callback():
    st.session_state["raw_text_input"] = ""
    st.session_state["process_triggered"] = False


def trigger_callback():
    st.session_state["process_triggered"] = True
    st.session_state["refresh_nonce"] += 1


with st.expander("Unos tabele", expanded=True):
    raw_text = st.text_area("Zalepite tabelu iz vašeg programa (Ctrl + V):", height=160, key="raw_text_input")
col1, col2, _ = st.columns([2, 1, 3])
col1.button("Učitaj i proveri", icon=":material/sync:", type="primary", use_container_width=True,
             on_click=trigger_callback, disabled=not operator_ok,
             help=None if operator_ok else "Prvo unesite ime operatera u levom meniju.")
col2.button("Očisti", icon=":material/delete:", on_click=clear_text_callback, use_container_width=True)


# ==========================================
# 10. REZULTATI
# ==========================================
cfg = dict(provider=provider, api_key=api_key, include_finished=include_finished, day=day, use_90=use_90,
           threshold=MATCH_THRESHOLD, grace=grace, auto=(refresh_mode == "Automatsko"), interval=auto_interval, sound_alert=sound_alert)
run_every = auto_interval if (refresh_mode == "Automatsko" and st.session_state["process_triggered"]) else None


@st.fragment(run_every=run_every)
def results_view(text, cfg):
    store = st.session_state["store"]
    inc_store, overrides = store["incidents"], store["overrides"]
    index = []

    if not st.session_state.get("operator", "").strip():
        st.warning("Unesite ime operatera u levom meniju. Bez imena se provera ne pokreće.")
    elif text and st.session_state["process_triggered"]:
        parsed = parse_copied_data(text)
        if not parsed:
            st.warning("Nije prepoznata struktura teksta. Proverite uslov kopiranja.")
        else:
            fixtures, remaining, limit, fetched_at, errors = [], "N/A", "N/A", None, []
            if cfg["api_key"]:
                fixtures, remaining, limit, fetched_at, errors = load_fixtures(
                    cfg["provider"], cfg["api_key"], cfg["include_finished"], cfg["day"], st.session_state["refresh_nonce"],
                    reuse=not cfg["auto"])
            else:
                st.warning("Nedostaje API ključ (secrets ili polje u sidebar-u).")
            for e in errors:
                st.error(f"{e}")
            if remaining != "N/A":
                st.session_state["api_remaining"], st.session_state["api_limit"] = remaining, limit

            index = build_index(fixtures, cfg["use_90"])
            by_id = {fx["id"]: fx for fx in index}
            pending, now = st.session_state["pending"], time.time()
            results, suggestions = [], []
            current_keys, ids_with_data, changed = set(), set(), False
            counts = {"ok": 0, "mismatch": 0, "wait": 0, "na": 0}

            for m in parsed:
                mid, my = m["ID"], m["Moj Sistem Rezultat"]
                ov = overrides.get(ov_key(cfg["provider"], mid))
                note, conf, fx = "", 0, None
                if ov:
                    fx, conf, note = by_id.get(ov["fixture_id"]), 100, "ručno"
                    if fx is None:
                        note = "ručno spojeno, nema u feedu"
                else:
                    fx, conf, amb = find_match(normalize_team_name(m["Home"]), normalize_team_name(m["Away"]),
                                               index, cfg["threshold"])
                    if amb:
                        fx, note = None, "dvosmisleno poklapanje"

                inc_state = ""
                if fx is None:
                    status, ext_score, ext_status = "Nema podataka", "N/A", note or "N/A"
                    counts["na"] += 1
                    pending.pop(mid, None)
                    cand = find_candidate(m, index)
                    if cand and not ov:
                        suggestions.append((m, cand))
                else:
                    ids_with_data.add(mid)
                    ext_score = fx["score"]
                    ext_status = fx["status"] + (f" ({note})" if note else "")
                    if my == "N/A":
                        status = "Sistem bez rezultata"
                        counts["na"] += 1
                    elif my == ext_score:
                        status = "OK"
                        counts["ok"] += 1
                        pending.pop(mid, None)
                    else:
                        key = f"{mid}|{my}|{ext_score}"
                        p = pending.get(mid)
                        if not p or p["key"] != key:
                            p = pending[mid] = {"key": key, "since": now}
                        left = cfg["grace"] - (now - p["since"])
                        if fx["short"] in FINISHED or left <= 0:
                            status = "Neslaganje"
                            counts["mismatch"] += 1
                            current_keys.add(key)
                            inc = inc_store.get(key)
                            if inc is None or inc["stanje"] == "rešeno":
                                inc_store[key] = {
                                    "id": mid, "liga": m["Liga"], "meč": m["Meč"], "sistem": my,
                                    "live": ext_score, "otkriveno": stamp(),
                                    "stanje": "aktivno", "operater": st.session_state.get("operator", "").strip(),
                                    "razreseno": "", "ts": now, "razreseno_ts": None, "alarmed": False}
                                changed = True
                            inc_state = inc_store[key]["stanje"]
                        else:
                            status = "Čeka potvrdu"
                            counts["wait"] += 1
                            ext_status += f" (još {int(left)}s)"

                results.append({
                    "Status": status, "Incident": inc_state, "ID": mid, "Oznaka": m["Oznaka"],
                    "Liga": m["Liga"], "Meč": m["Meč"], "Tvoj Sistem": my,
                    "Score2 (HT)": m["Score2 (HT)"], "Teren / Live Feed": ext_score,
                    "Status Meča": ext_status, "Pouzdanost": int(conf)})

            for key, inc in inc_store.items():
                if inc["stanje"] != "rešeno" and inc["id"] in ids_with_data and key not in current_keys:
                    inc["stanje"], inc["razreseno"], inc["razreseno_ts"] = "rešeno", stamp(), now
                    changed = True

            fresh = [i for i in inc_store.values() if i["stanje"] != "rešeno" and not i.get("alarmed")]
            if fresh and cfg["sound_alert"]:
                st.audio(make_wav(), format="audio/wav", autoplay=True)
            for i in fresh:
                i["alarmed"], changed = True, True
            if changed:
                save_store()

            active = {k: v for k, v in inc_store.items() if v["stanje"] != "rešeno"}
            age = (now - fetched_at) if fetched_at else None
            problems = []
            if not cfg["api_key"]:
                problems.append("nema API ključa")
            elif errors:
                problems.append("greška pri učitavanju feeda")
            elif not index:
                problems.append("API nije vratio nijednu utakmicu")
            if len(parsed) >= 5 and counts["na"] / len(parsed) >= 0.8:
                problems.append("većina mečeva nema podataka")
            if age is not None and cfg["auto"] and age > max(90, 3 * cfg["interval"]):
                problems.append(f"podaci stariji od {int(age)} s")
            feed_problem = "; ".join(problems)
            status_banner(len(active), counts["wait"], counts["na"], counts["ok"], feed_problem)
            st.caption(f"Podaci iz feeda: {fmt_time(fetched_at)} (pre {int(age)} s). Poslednja provera: {fmt_time(now)}."
                       if fetched_at else f"Podaci iz feeda nisu učitani. Poslednja provera: {fmt_time(now)}.")
            if active:
                with st.container(border=True):
                    for v in active.values():
                        st.markdown(f"**ID {v['id']}**, {v['meč']}: sistem **{v['sistem']}** → "
                                    f"live :red[**{v['live']}**], {v['otkriveno']}, operater: {v['operater'] or '—'}")

            df_full = pd.DataFrame(results)
            df_full["Rezultat"] = df_full["Tvoj Sistem"] + " → " + df_full["Teren / Live Feed"]
            df_full["_r"] = df_full["Status"].map(STATUS_RANK)
            df_full = df_full.sort_values("_r", kind="stable").drop(columns="_r")
            m1, m2, m3, m4, m5, m6 = st.columns(6)
            m1.metric("Ukupno mečeva", len(df_full))
            m2.metric("Usklađeno (OK)", counts["ok"])
            m3.metric("Neslaganja", counts["mismatch"], delta_color="inverse")
            m4.metric("Čeka potvrdu", counts["wait"])
            m5.metric("Nema podataka", counts["na"])
            m6.metric("API zahteva preostalo", st.session_state.get("api_remaining", "N/A"))

            rem, lim = st.session_state.get("api_remaining"), st.session_state.get("api_limit")
            if str(rem).isdigit() and str(lim).isdigit() and int(lim) > 0:
                used_frac = 1 - int(rem) / int(lim)
                st.progress(min(max(used_frac, 0.0), 1.0),
                            text=f"API: potrošeno {int(lim) - int(rem)} od {lim} dnevnih zahteva (preostalo {rem})")

            q = st.text_input("Pretraga", "", key="search_q", placeholder="ID, klub ili liga")
            dq = df_full
            if q:
                dq = dq[dq["ID"].astype(str).str.contains(q, case=False, regex=False)
                        | dq["Meč"].str.contains(q, case=False, regex=False)
                        | dq["Liga"].str.contains(q, case=False, regex=False)]
            groups = [
                ("Neslaganja", dq[dq["Status"] == "Neslaganje"]),
                ("Čeka potvrdu", dq[dq["Status"] == "Čeka potvrdu"]),
                ("N/A", dq[dq["Status"].isin(["Nema podataka", "Sistem bez rezultata"])]),
                ("Sve", dq),
            ]
            for tab, (name, d) in zip(st.tabs([f"{n} ({len(d)})" for n, d in groups]), groups):
                with tab:
                    render_table(d)

            with st.expander(f"Ručno spajanje ({len(overrides)} sačuvanih)"):
                if index:
                    mby = {m["ID"]: m for m in parsed}
                    sel_m = st.selectbox("Meč iz tvog sistema:", [m["ID"] for m in parsed],
                                         format_func=lambda i: f"{i} – {mby[i]['Meč']}", key="ov_match")
                    sel_f = st.selectbox("Utakmica na API-ju:", [f["id"] for f in index],
                                         format_func=lambda i: f"{by_id[i]['home']} - {by_id[i]['away']} "
                                                               f"({by_id[i]['liga']}, {by_id[i]['status']})",
                                         key="ov_fix")
                    st.button("Sačuvaj spajanje", on_click=lambda: apply_override(mby[sel_m], by_id[sel_f]))
                else:
                    st.info("Nema učitanih utakmica sa API-ja.")
                if overrides:
                    st.button("Obriši sva ručna spajanja", on_click=clear_overrides)

            if suggestions:
                with st.expander("Predlozi za spajanje (jedan klik)", expanded=True):
                    for m, cand in suggestions:
                        c1, c2, c3 = st.columns([2, 3, 2])
                        c1.write(f"**ID {m['ID']}**: {m['Meč']}")
                        c2.write(f"API kandidat: **{cand['home']} - {cand['away']}** ({cand['status']})")
                        c3.button("Spoji", key=f"sug_{m['ID']}", on_click=apply_override, args=(m, cand))

            meta = [
                ("Generisano", datetime.now(TZ).strftime("%d.%m.%Y %H:%M:%S")),
                ("Operater", st.session_state.get("operator", "").strip() or "—"),
                ("Provajder", cfg["provider"]),
                ("Podaci iz feeda", fmt_time(fetched_at) if fetched_at else "—"),
                ("Izvor", f"uživo + završeni ({cfg['day']})" if cfg["include_finished"] else "samo uživo"),
                ("API zahteva preostalo", st.session_state.get("api_remaining", "N/A")),
                ("Ukupno mečeva", len(df_full)), ("Neslaganja", counts["mismatch"]),
                ("Čeka potvrdu", counts["wait"]), ("Bez podataka", counts["na"]),
                ("Upozorenje feeda", feed_problem or "nema"),
            ]
            d1, d2, _ = st.columns([2, 2, 3])
            d1.download_button("Preuzmi Excel izveštaj", build_excel(df_full, inc_store, meta),
                               file_name=f"settlementcheck_{datetime.now(TZ).strftime('%Y%m%d_%H%M')}.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               icon=":material/download:")
            d2.download_button("Preuzmi CSV", df_full.to_csv(index=False).encode("utf-8"),
                               file_name="settlementcheck_report.csv", mime="text/csv")

    if inc_store:
        st.markdown("---")
        st.subheader("Arhiva neslaganja")
        st.caption("Pamti se i posle osvežavanja stranice (settlement_state.json).")
        arch = incidents_frame(inc_store, time.time())
        st.dataframe(arch.iloc[::-1], use_container_width=True, hide_index=True)
        st.button("Obriši rešena neslaganja", on_click=clear_resolved)

    if index:
        with st.expander("Sve utakmice trenutno dostupne sa API-ja"):
            st.dataframe(pd.DataFrame([{"Liga": f["liga"], "Domaćin": f["home"], "Gost": f["away"],
                                        "Rezultat": f["score"], "Status": f["status"]} for f in index]),
                         use_container_width=True, hide_index=True)


results_view(raw_text, cfg)
